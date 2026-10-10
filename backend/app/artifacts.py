from __future__ import annotations

import io
import re
from dataclasses import dataclass

CANVAS_RE = re.compile(r"<canvas\b([^>]*)>(.*?)</canvas>", re.IGNORECASE | re.DOTALL)
ATTR_RE = re.compile(r'(\w+)\s*=\s*"([^"]*)"')

CHARS_PER_TOKEN = 3.6
VALID_KINDS = {"markdown", "code", "text", "html", "svg", "json", "csv", "mermaid"}
_KIND_ALIASES = {"md": "markdown", "htm": "html", "htmx": "html", "js": "code", "py": "code", "xml": "html"}


def normalize_kind(kind: str) -> str:
    kind = (kind or "markdown").strip().lower()
    kind = _KIND_ALIASES.get(kind, kind)
    return kind if kind in VALID_KINDS else "markdown"


@dataclass
class ParsedCanvas:
    title: str
    kind: str
    language: str
    content: str


def parse_canvas_blocks(text: str) -> tuple[str, list[ParsedCanvas]]:
    """Split a model reply into (chat_text, canvas artifacts).

    Canvas content is removed from the chat text and replaced by a short
    placeholder so the stored (and re-sent) message stays small.
    """
    blocks: list[ParsedCanvas] = []

    def _replace(match: re.Match[str]) -> str:
        attrs = dict(ATTR_RE.findall(match.group(1)))
        content = match.group(2).strip("\n")
        kind = normalize_kind(attrs.get("kind", "markdown"))
        title = attrs.get("title", "").strip() or _guess_title(content, kind)
        blocks.append(
            ParsedCanvas(
                title=title,
                kind=kind,
                language=attrs.get("language", "").strip(),
                content=content,
            )
        )
        return f"\n\n[canvas: {title}]\n\n"

    cleaned = CANVAS_RE.sub(_replace, text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    return cleaned, blocks


def _guess_title(content: str, kind: str) -> str:
    for line in content.splitlines():
        line = line.strip().lstrip("#").strip()
        if line:
            return (line[:80] or f"{kind} artifact").strip()
    return f"{kind} artifact"


def tokens_of(text: str) -> int:
    return max(1, int(len(text) / CHARS_PER_TOKEN))


# ------------------------------------------------------------------ exports
def to_markdown(title: str, content: str) -> str:
    return content if content.lstrip().startswith("#") else f"# {title}\n\n{content}"


def to_text(title: str, content: str) -> str:
    text = re.sub(r"^#{1,6}\s+", "", content, flags=re.MULTILINE)
    text = re.sub(r"[*_`>]", "", text)
    return text


def to_html(title: str, content: str, kind: str = "markdown") -> str:
    import html as html_mod

    kind = normalize_kind(kind)
    title_esc = html_mod.escape(title)

    # Already a full HTML document.
    stripped = content.lstrip().lower()
    if kind == "html" and ("<!doctype" in stripped[:64] or "<html" in stripped[:64]):
        return content

    body_style = (
        "font-family:Georgia,'Times New Roman',serif;max-width:820px;"
        "margin:40px auto;padding:0 20px;line-height:1.7;color:#111"
    )

    if kind == "html":
        body = content
    elif kind == "svg":
        body = content
    else:
        body = f"<pre style='white-space:pre-wrap;font-family:monospace'>{html_mod.escape(content)}</pre>"

    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>{title_esc}</title>"
        "<style>img,svg{max-width:100%}table{border-collapse:collapse}"
        "td,th{border:1px solid #ccc;padding:6px 10px}</style>"
        f'</head><body style="{body_style}">{body}</body></html>'
    )


def to_svg(content: str) -> str:
    return content


def to_json(content: str) -> str:
    import json as json_mod

    try:
        return json_mod.dumps(json_mod.loads(content), indent=2, ensure_ascii=False)
    except (ValueError, TypeError):
        return content


def _inline_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


def to_docx(title: str, content: str) -> bytes:
    """Render Markdown-ish content to a .docx using python-docx."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=0)

    in_code = False
    for raw in content.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph(line, style="No Spacing")
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            doc.add_heading(_inline_md(heading.group(2)), level=min(len(heading.group(1)), 6))
            continue
        if re.match(r"^\s*[-*+]\s+", line):
            doc.add_paragraph(_inline_md(re.sub(r"^\s*[-*+]\s+", "", line)), style="List Bullet")
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            doc.add_paragraph(_inline_md(re.sub(r"^\s*\d+\.\s+", "", line)), style="List Number")
            continue
        if line.strip().startswith(">"):
            doc.add_paragraph(_inline_md(line.strip().lstrip("> ")), style="Intense Quote")
            continue
        doc.add_paragraph(_inline_md(line))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
